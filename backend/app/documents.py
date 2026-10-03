"""Resume text extraction and offer-letter PDF rendering."""
import io
from datetime import date

from docx import Document
from fpdf import FPDF
from pypdf import PdfReader

from .config import settings

ALLOWED_EXT = {".pdf", ".docx", ".txt"}


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    return data.decode("utf-8", errors="ignore").strip()


def _latin(text: str) -> str:
    """Core PDF fonts are Latin-1 only; replace common Unicode punctuation."""
    table = {"₹": "INR ", "–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "•": "-", "…": "..."}
    for k, v in table.items():
        text = text.replace(k, v)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def offer_letter_pdf(offer: dict, candidate_name: str) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_margins(20, 20, 20)

    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(79, 70, 229)
    pdf.cell(0, 12, _latin(settings.COMPANY_NAME), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 6, f"Date: {date.today().strftime('%d %B %Y')}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    pdf.set_text_color(15, 23, 42)
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, "Offer of Employment", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 8, _latin(f"Dear {candidate_name},"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.multi_cell(0, 6, _latin(offer.get("letter_body", "")), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Key Terms", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    rows = [
        ("Position", offer.get("position", "")),
        ("Department", offer.get("department", "")),
        ("Location", offer.get("location", "")),
        ("Annual CTC", f"{offer.get('currency', 'INR')} {offer.get('annual_ctc', 0):,.0f}"),
        ("Joining date", offer.get("joining_date", "")),
        ("Probation", offer.get("probation", "")),
        ("Offer valid until", offer.get("valid_until", "")),
    ]
    for label, value in rows:
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(45, 7, _latin(label))
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 7, _latin(str(value)), new_x="LMARGIN", new_y="NEXT")
    if offer.get("benefits"):
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "Benefits", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for b in offer["benefits"]:
            pdf.multi_cell(0, 6, _latin(f"- {b}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(8)
    pdf.multi_cell(0, 6, _latin("Please accept or decline this offer using the secure link shared with you by email. "
                                "This offer is subject to background verification as per company policy."),
                   new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "For " + _latin(settings.COMPANY_NAME), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 6, "Talent Acquisition Team", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
